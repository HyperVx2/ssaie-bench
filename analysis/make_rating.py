import json, random, collections
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
L="ABCD"
def rows(p): return [json.loads(l) for l in open(p, encoding="utf-8")]
lap={ (r["model"],r["id"]):r for r in rows("results/laptop/quality.jsonl") if r["task"]=="B" and r["status"]=="ok"}
gen={ ("gemini",r["id"]):r for r in rows("results/frontier/generated.jsonl")}
topics={t["id"]:t for t in json.load(open("data/topics.json"))}
J=[r for r in rows("results/frontier/judgements.jsonl") if r.get("judged")]
pool=collections.defaultdict(list)
for j in J: pool[j["model"]].append(j)
rng=random.Random(2027)
picked=[]
for m in sorted(pool):
    picked += rng.sample(pool[m], min(4,len(pool[m])))
rest=[j for j in J if j not in picked]
picked += rng.sample(rest, 40-len(picked))
rng.shuffle(picked)
def item(j):
    src=(lap if j["model"]!="gemini" else gen)[(j["model"],j["id"])]
    s=src["response"]; obj=json.loads(s[s.find("{"):s.rfind("}")+1])
    opts=obj["options"]
    if isinstance(opts,list): opts={L[i]:v for i,v in enumerate(opts)}
    new={L[i]:str(opts[o]) for i,o in enumerate(j["order"])}
    return obj["stem"], new, j["shown_key"]
key=[]
wb=Workbook()
ins=wb.active; ins.title="Instructions"
F=lambda **k: Font(name="Arial", **k)
lines=[("Blind rating sheet — 40 generated exam items",True),
("",False),
("What to do",True),
("1. Go to the 'Rating' tab. Fill in only the yellow columns (J–N) for all 40 items.",False),
("2. Judge each item as you would an item a colleague wrote for a 1st/2nd-year CS exam. Do not look anything up.",False),
("3. You are not told which model wrote each item. Please don't try to guess.",False),
("4. Save the file and send it back. Expected time: about 30 minutes.",False),
("",False),
("Rubric",True),
("Key correct? (Y/N): Y if the marked key is the single correct answer. N if it is wrong, or if more than one option is correct.",False),
("Clarity (1–4): 1 = confusing or ambiguous · 2 = understandable but flawed · 3 = clear with minor issues · 4 = perfectly clear.",False),
("Distractors (1–4): 1 = wrong options are implausible or obviously wrong · 4 = all wrong options are plausible to a weak student.",False),
("Matches topic & Bloom? (Y/N): Y if the item tests the stated topic at roughly the stated Bloom level.",False),
("Notes: optional, a few words (e.g., 'two correct options', 'answer is C').",False),
("",False),
("Example of a filled row (not one of the 40 items)",True)]
for i,(t,b) in enumerate(lines,1):
    c=ins.cell(row=i,column=1,value=t); c.font=F(bold=b, size=13 if i==1 else 11)
hdr=["#","Topic","Bloom level","Question","A","B","C","D","Marked key","Key correct? (Y/N)","Clarity (1–4)","Distractors (1–4)","Matches topic & Bloom? (Y/N)","Notes"]
ex=["Ex","Stacks and queues: order of operations","Apply","Push 1, 2, 3 onto an empty stack, then pop once. Which element is on top?","1","2","3","The stack is empty","B","Y",4,3,"Y","clear; D is weak"]
r0=len(lines)+1
thin=Side(style="thin",color="BBBBBB")
for c,v in enumerate(hdr,1):
    x=ins.cell(row=r0,column=c,value=v); x.font=F(bold=True); x.fill=PatternFill("solid",fgColor="DDDDDD")
for c,v in enumerate(ex,1):
    x=ins.cell(row=r0+1,column=c,value=v); x.font=F(); x.alignment=Alignment(wrap_text=True,vertical="top")
ins.column_dimensions["A"].width=110
ws=wb.create_sheet("Rating")
for c,v in enumerate(hdr,1):
    x=ws.cell(row=1,column=c,value=v); x.font=F(bold=True,color="FFFFFF"); x.fill=PatternFill("solid",fgColor="404040")
    x.alignment=Alignment(wrap_text=True,vertical="center")
yellow=PatternFill("solid",fgColor="FFFF00")
for n,j in enumerate(picked,1):
    stem,opts,k=item(j); t=topics[j["id"]]
    vals=[n,t["topic"],t["bloom"],stem,opts["A"],opts["B"],opts["C"],opts["D"],k,None,None,None,None,None]
    for c,v in enumerate(vals,1):
        x=ws.cell(row=n+1,column=c,value=v); x.font=F(); x.alignment=Alignment(wrap_text=True,vertical="top")
        x.border=Border(top=thin,bottom=thin,left=thin,right=thin)
        if c>=10: x.fill=yellow
    key.append({"n":n,"uid":j["uid"],"model":j["model"],"id":j["id"],"shown_key":k,
                "judge":{q:j.get(q) for q in ("judge_answer","key_correct","clarity","distractors","aligned")}})
widths=[5,26,11,55,22,22,22,22,8,11,10,11,13,22]
for i,w in enumerate(widths,1): ws.column_dimensions[ws.cell(row=1,column=i).column_letter].width=w
ws.freeze_panes="B2"
yn=DataValidation(type="list",formula1='"Y,N"',allow_blank=True); s14=DataValidation(type="list",formula1='"1,2,3,4"',allow_blank=True)
ws.add_data_validation(yn); ws.add_data_validation(s14)
yn.add("J2:J41"); yn.add("M2:M41"); s14.add("K2:K41"); s14.add("L2:L41")
ws.cell(row=43,column=2,value="Rows completed (all four ratings filled):").font=F(bold=True)
ws.cell(row=43,column=4,value='=SUMPRODUCT((J2:J41<>"")*(K2:K41<>"")*(L2:L41<>"")*(M2:M41<>""))').font=F(bold=True)
ws.cell(row=44,column=2,value="Target: 40").font=F(italic=True)
wb.save("/mnt/user-data/outputs/SSAIE_rating_sheet.xlsx")
json.dump(key,open("rating_key_HIDDEN.json","w"),indent=1)
print(collections.Counter(k["model"] for k in key))
